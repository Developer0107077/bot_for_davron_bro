import io
import re
from aiogram import Bot, BaseMiddleware
from aiogram.enums import ChatMemberStatus
from aiogram.exceptions import TelegramBadRequest
from aiogram.types import CallbackQuery, Message
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config
import db


CODE_RE = re.compile(r"^[A-Z0-9_\-]{1,20}$")


def norm_code(s) -> str:
    """'r 3' -> 'R3' (bo'shliqlarsiz, katta harf)."""
    return re.sub(r"\s+", "", s or "").upper()


async def show(ev, text, markup=None):
    """Callback bo'lsa xabarni tahrirlaydi, aks holda yangi xabar yuboradi."""
    if isinstance(ev, CallbackQuery):
        try:
            await ev.message.edit_text(text, reply_markup=markup)
        except TelegramBadRequest as e:
            if "not modified" not in str(e):
                await ev.message.answer(text, reply_markup=markup)
        try:
            await ev.answer()
        except Exception:
            pass
    else:
        await ev.answer(text, reply_markup=markup)


async def check_sub(bot: Bot, uid: int) -> bool:
    if not config.CHANNEL_ID or uid in config.ADMIN_IDS:
        return True
    try:
        m = await bot.get_chat_member(config.CHANNEL_ID, uid)
        return m.status not in (ChatMemberStatus.LEFT, ChatMemberStatus.KICKED)
    except Exception:
        return True  # kanal sozlanmagan yoki bot admin emas — to'smaymiz


def make_xlsx(sheet, headers, rows) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F75B5")
        c.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(list(r))
    for i, col in enumerate(ws.columns, 1):
        w = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(w + 3, 50)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class BlockMiddleware(BaseMiddleware):
    """Bloklangan foydalanuvchilarni to'xtatadi (adminlar bundan mustasno)."""

    async def __call__(self, handler, event, data):
        u = event.from_user
        if u and u.id not in config.ADMIN_IDS:
            row = await db.fetchone("SELECT is_blocked FROM users WHERE tg_id=?", (u.id,))
            if row and row["is_blocked"]:
                if isinstance(event, Message):
                    await event.answer("🚫 Siz bloklangansiz. Murojaat uchun admin bilan bog'laning.")
                else:
                    await event.answer("🚫 Siz bloklangansiz.", show_alert=True)
                return
        return await handler(event, data)
